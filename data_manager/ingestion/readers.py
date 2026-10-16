"""
Streaming, bounded-memory readers for CSV, XLSX and ODS.

Every reader yields rows as lists of strings, so geographic codes, leading zeros and labels are
never re-typed. Numeric spreadsheet cells are returned as their stored value (office:value for
ODS, the cell value for XLSX), not as display text, so "8,235" arrives as "8235".

ODS is parsed with a streaming XML parser over content.xml inside the zip: only the requested
sheet is materialised, one row at a time, and each processed row is detached from the tree.
pandas' odf engine (odfpy) builds a full DOM; on the real VEH0142 workbook it exceeded 2.6 GB
before failing, so it is not used.
"""

import codecs
import csv
import html
import io
import os
import re
import zipfile

from .archive_safety import UnsafeArchive, inspect_archive

try:  # defusedxml refuses DTDs and entity declarations; fall back to an explicit pre-scan.
    from defusedxml.ElementTree import iterparse as _safe_iterparse

    def _iterparse(f):
        return _safe_iterparse(f, events=("start", "end"), forbid_dtd=True)
except ImportError:  # pragma: no cover - defusedxml is present in this environment
    from xml.etree.ElementTree import iterparse as _plain_iterparse

    def _iterparse(f):
        return _plain_iterparse(f, events=("start", "end"))

T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
O = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
X = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
NUMERIC_TYPES = ("float", "percentage", "currency")


class UnsupportedFormat(ValueError):
    """The bytes are not CSV, XLSX or ODS."""


class ParseError(ValueError):
    """The file could not be parsed safely."""


# ------------------------------------------------------------------ format detection
def sniff_format(path, original_name=""):
    """Decide the format from the bytes; the extension must agree."""
    with open(path, "rb") as f:
        head = f.read(65536)
    ext = original_name.lower().rsplit(".", 1)[-1] if "." in original_name else ""
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as z:
                names = set(z.namelist())
                mimetype = z.read("mimetype").strip() if "mimetype" in names else b""
        except zipfile.BadZipFile as e:
            raise UnsupportedFormat(f"corrupt zip container: {e}") from e
        if mimetype == b"application/vnd.oasis.opendocument.spreadsheet":
            fmt = "ods"
        elif "xl/workbook.xml" in names:
            fmt = "xlsx"
        else:
            raise UnsupportedFormat("zip file is neither an ODS nor an XLSX workbook")
    elif head.startswith(b"\xd0\xcf\x11\xe0"):
        raise UnsupportedFormat("legacy binary Excel (.xls/.xlsb) is not supported; save as .xlsx, .ods or .csv")
    elif b"\x00" in head:
        raise UnsupportedFormat("binary content is not a CSV file")
    else:
        fmt = "csv"
    if ext and ext != fmt:
        raise UnsupportedFormat(f"file extension .{ext} does not match its content ({fmt.upper()})")
    return fmt


# ------------------------------------------------------------------ progress
class _CountingReader(io.RawIOBase):
    """Wraps a binary stream; reports bytes consumed to a callback (throttled by the caller)."""

    def __init__(self, raw, callback=None, prefix=b""):
        self.raw, self.callback, self.count, self.prefix = raw, callback, 0, prefix

    def readable(self):
        return True

    def readinto(self, b):
        if self.prefix:
            data, self.prefix = self.prefix[:len(b)], self.prefix[len(b):]
        else:
            data = self.raw.read(len(b))
        n = len(data)
        b[:n] = data
        self.count += n
        if self.callback and n:
            self.callback(self.count)
        return n


# ------------------------------------------------------------------ CSV
def _csv_encoding(path):
    """utf-8 (with or without BOM) if the whole file decodes, otherwise cp1252."""
    dec = codecs.getincrementaldecoder("utf-8")()
    with open(path, "rb") as f:
        try:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                dec.decode(chunk)
            dec.decode(b"", final=True)
            return "utf-8-sig"
        except UnicodeDecodeError:
            return "cp1252"


def iter_csv_rows(path, progress=None):
    encoding = _csv_encoding(path)
    with open(path, "rb") as raw:
        counted = io.BufferedReader(_CountingReader(raw, progress), buffer_size=1 << 20)
        text = io.TextIOWrapper(counted, encoding=encoding, newline="")
        try:
            for row in csv.reader(text):
                yield row
        except csv.Error as e:
            raise ParseError(f"CSV syntax error: {e}") from e


# ------------------------------------------------------------------ XLSX
def xlsx_sheet_names(path, settings):
    inspect_archive(path, settings, "xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _xlsx_cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def iter_xlsx_rows(path, sheet, settings, progress=None):
    inspect_archive(path, settings, "xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ParseError(f"worksheet {sheet!r} not found (sheets: {wb.sheetnames})")
        ws = wb[sheet]
        total = ws.max_row or 0
        for i, values in enumerate(ws.iter_rows(values_only=True), 1):
            cells = [_xlsx_cell(v) for v in values]
            while cells and cells[-1] == "":
                cells.pop()
            if len(cells) > settings.max_columns:
                raise ParseError(f"row {i} has {len(cells)} columns (limit {settings.max_columns})")
            if progress and total and i % 2000 == 0:
                progress(i / total)
            yield cells
    finally:
        wb.close()


# ------------------------------------------------------------------ ODS
def _ods_text(cell):
    """Text content of an ODS cell: paragraphs joined by newlines; text:s / text:tab / line breaks expanded."""
    parts = []

    def walk(el):
        if el.text:
            parts.append(el.text)
        for ch in el:
            tag = ch.tag
            if tag == X + "s":
                parts.append(" " * int(ch.get(X + "c", "1")))
            elif tag == X + "tab":
                parts.append("\t")
            elif tag == X + "line-break":
                parts.append("\n")
            elif tag != T + "annotation":
                walk(ch)
            if ch.tail:
                parts.append(ch.tail)

    paragraphs = []
    for p in cell:
        if p.tag in (X + "p", X + "h"):
            parts = []
            walk(p)
            paragraphs.append("".join(parts))
    return "\n".join(paragraphs)


def normalise_number(text):
    """Stored spreadsheet number -> shortest exact decimal text ('0.42599999999999999' -> '0.426', '8235.0' -> '8235')."""
    try:
        v = float(text)
    except ValueError:
        return text
    if v != v or v in (float("inf"), float("-inf")):
        return text
    if v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return repr(v)


def _ods_value(cell):
    vt = cell.get(O + "value-type")
    if vt in NUMERIC_TYPES:
        return normalise_number(cell.get(O + "value", ""))
    if vt == "date":
        return cell.get(O + "date-value", "")
    if vt == "time":
        return cell.get(O + "time-value", "")
    if vt == "boolean":
        return cell.get(O + "boolean-value", "")
    return _ods_text(cell)


def _ods_row(row, settings, index):
    cells, pending = [], 0
    for c in row:
        if c.tag not in (T + "table-cell", T + "covered-table-cell"):
            continue
        reps = int(c.get(T + "number-columns-repeated", "1"))
        value = _ods_value(c)
        if value == "":
            pending += reps
            continue
        if len(cells) + pending + reps > settings.max_columns:
            raise ParseError(f"row {index} exceeds {settings.max_columns} columns")
        cells.extend([""] * pending)
        pending = 0
        cells.extend([value] * reps)
    return cells


def _open_content(path, settings, progress):
    inspect_archive(path, settings, "ods")
    z = zipfile.ZipFile(path)
    info = z.getinfo("content.xml")
    raw = z.open(info)
    try:
        head = raw.read(65536)
    except zipfile.BadZipFile as e:
        raw.close()
        z.close()
        raise UnsafeArchive(f"archive member failed its integrity check: {e}") from e
    if b"<!DOCTYPE" in head or b"<!ENTITY" in head:
        raw.close()
        z.close()
        raise UnsafeArchive("content.xml declares a DTD or entities; refused")
    total = info.file_size

    def cb(n):
        if progress:
            progress(n / total if total else 0.0)
    return z, _CountingReader(raw, cb, prefix=head)


_TABLE_TAG = re.compile(rb'<table:table\s[^>]*?table:name="([^"]*)"')


def ods_sheet_names_fast(path, settings, progress=None):
    """List sheets by scanning the decompressed bytes for table start tags (no XML parse).

    '<' cannot occur inside a well-formed tag, so every tag that starts before the last '<' of a
    chunk is complete; the remainder is carried into the next chunk. Falls back to the XML
    parser if the conventional 'table:' prefix is not used.
    """
    z, f = _open_content(path, settings, progress)
    names, tail = [], b""
    try:
        while True:
            chunk = f.read(4 << 20)
            if not chunk:
                break
            buf = tail + chunk
            cut = buf.rfind(b"<")
            cut = cut if cut >= 0 else len(buf)
            for m in _TABLE_TAG.finditer(buf, 0, cut):
                names.append(html.unescape(m.group(1).decode("utf-8")))
            tail = buf[cut:]
        for m in _TABLE_TAG.finditer(tail):
            names.append(html.unescape(m.group(1).decode("utf-8")))
    except zipfile.BadZipFile as e:
        raise UnsafeArchive(f"archive member failed its integrity check: {e}") from e
    finally:
        z.close()
    return names or ods_sheet_names(path, settings)


def ods_sheet_names(path, settings, progress=None):
    """List sheets by streaming content.xml (no DOM)."""
    z, f = _open_content(path, settings, progress)
    names = []
    try:
        for ev, el in _iterparse(f):
            if ev == "start" and el.tag == T + "table":
                names.append(el.get(T + "name"))
            elif ev == "end" and el.tag == T + "table-row":
                el.clear()
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"ODS content.xml is malformed: {type(e).__name__}: {e}") from e
    finally:
        z.close()
    return names


def iter_ods_rows(path, sheet, settings, progress=None):
    """Yield the rows of one sheet. Blank rows are yielded as [] (repeated blank runs collapse to one)."""
    z, f = _open_content(path, settings, progress)
    stack, in_sheet, found, index = [], False, False, 0
    try:
        try:
            for ev, el in _iterparse(f):
                if ev == "start":
                    stack.append(el)
                    if el.tag == T + "table":
                        in_sheet = el.get(T + "name") == sheet
                        found = found or in_sheet
                    continue
                stack.pop()
                if el.tag == T + "table-row":
                    if in_sheet:
                        reps = int(el.get(T + "number-rows-repeated", "1"))
                        cells = _ods_row(el, settings, index + 1)
                        if cells and reps > settings.max_repeated_rows:
                            raise ParseError(f"row {index + 1} repeated {reps} times (limit {settings.max_repeated_rows})")
                        for _ in range(reps if cells else 1):
                            index += 1
                            yield cells
                        if not cells:
                            index += reps - 1
                    el.clear()
                    if stack:
                        stack[-1].remove(el)
                elif el.tag == T + "table":
                    el.clear()
                    if in_sheet:
                        return
        except (ParseError, UnsafeArchive):
            raise
        except zipfile.BadZipFile as e:
            raise UnsafeArchive(f"archive member failed its integrity check: {e}") from e
        except Exception as e:  # noqa: BLE001 - expat/defusedxml errors
            raise ParseError(f"ODS content.xml is malformed: {type(e).__name__}: {e}") from e
        if not found:
            raise ParseError(f"worksheet {sheet!r} not found")
    finally:
        z.close()


# ------------------------------------------------------------------ dispatch
def sheet_names(path, fmt, settings, progress=None):
    if fmt == "csv":
        return [None]
    if fmt == "xlsx":
        return xlsx_sheet_names(path, settings)
    if fmt == "ods":
        return ods_sheet_names_fast(path, settings, progress)
    raise UnsupportedFormat(fmt)


def iter_rows(path, fmt, sheet, settings, progress=None):
    if fmt == "csv":
        total = os.path.getsize(path)
        return iter_csv_rows(path, (lambda n: progress(n / total)) if progress and total else None)
    if fmt == "xlsx":
        return iter_xlsx_rows(path, sheet, settings, progress)
    if fmt == "ods":
        return iter_ods_rows(path, sheet, settings, progress)
    raise UnsupportedFormat(fmt)
